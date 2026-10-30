"""Whatever is thrown at the API, it answers with a clean 4xx or a valid result, never a 500."""

import io
import re

import pytest

from app.main import app
from app.security import public_track_limiter

NASTY = ["'; DROP TABLE instruments; --", "\" OR 1=1 --", "<script>alert(1)</script>", "कदम्ब प्रौद्योगिकी संस्थान",
         "🏫🔥💸", "‮RTL override", "a" * 5000, "%00%ff", "../../etc/passwd", "-1", "99999999999999999999999", "", " "]
BAD_IDS = ["999999", "0", "-1", "abc", "99999999999999999999999", "1.5", "%27"]
SKIP = {"/api/stream", "/api/auth/login"}


def _routes():
    paths = app.openapi()["paths"]
    assert len(paths) > 40, "route discovery broke"
    for path, ops in paths.items():
        if path.startswith("/api") and path not in SKIP:
            for m in ops:
                if m in ("get", "post", "put", "patch", "delete"):
                    yield m.upper(), path


@pytest.mark.parametrize("handle", ["rb.officer1", "auditor", "kadamb"])
def test_every_route_survives_bad_ids_and_garbage_bodies(api, handle):
    headers = api.h(handle)
    failures = []
    for method, path in _routes():
        for bad in (BAD_IDS if "{" in path else [""]):
            url = re.sub(r"\{[^}]+\}", bad, path)
            bodies = [None] if method == "GET" else [{}, [], "x", {"remarks": NASTY[0], "officer_id": "x",
                                                                   "amount": NASTY[3], "checklist": "x"}]
            for body in bodies:
                public_track_limiter.reset()
                try:
                    r = api.client.request(method, url, headers=headers, json=body)
                except Exception as e:  # noqa: BLE001 - the test client re-raises server errors
                    failures.append((method, url, body, repr(e)[:120]))
                    continue
                if r.status_code >= 500:
                    failures.append((method, url, body, r.status_code))
    assert not failures, failures[:10]


def test_query_strings_with_injection_are_harmless(api):
    before = api.get("rb.head", "/instruments").json()["total"]
    for q in NASTY:
        for path, params in [("/instruments", {"q": q}), ("/instruments", {"sort": q, "order": q}),
                             ("/instruments", {"status": q}), ("/instruments", {"expiring_days": q}),
                             ("/instruments", {"page_size": q}), ("/institutions", {"q": q, "sort": q}),
                             ("/alerts", {"kind": q}), ("/admin/audit", {"action": q}),
                             ("/reports/maturity", {"date_from": q}), ("/dashboard/summary", {"region_id": q})]:
            r = api.get("rb.head", path, params=params)
            assert r.status_code < 500, (path, params, r.status_code, r.text[:200])
        r = api.post("rb.head", "/ask", {"question": q or "x"})
        assert r.status_code < 500
    assert api.get("rb.head", "/instruments").json()["total"] == before


def test_public_letter_verification_with_garbage(api):
    for q in NASTY + ["NDL-00000000", "ndl-'or'1"]:
        public_track_limiter.reset()
        r = api.client.get(f"/api/public/verify/{q}")
        assert r.status_code in (404, 422, 429, 405), (q[:30], r.status_code)


@pytest.mark.parametrize("patch", [
    {"issue_date": "2026-02-30"}, {"issue_date": "yesterday"}, {"issue_date": "1900-01-01"},
    {"maturity_date": "2020-01-01"}, {"amount": -5}, {"amount": "lots"}, {"amount": 1e15}, {"bank_id": 999999},
    {"purpose": "BRIBE"}, {"fdr_number": ""}, {"depositor_name": None}, {"approval_year": "2026"},
    {"replaces_id": 999999}, {"ifsc": "x" * 40}, {"institution_id": 999999},
])
def test_invalid_submission_fields_are_rejected_cleanly(api, patch):
    r = api.post("rb.officer1", "/instruments", api.body(**patch))
    assert r.status_code in (400, 403, 404, 422), (patch, r.status_code, r.text[:300])


@pytest.mark.parametrize("text", ["कदम्ब शिक्षण प्रसारक मंडळ", "Trust 🙏 Society", "'; DELETE FROM users; --"])
def test_unicode_and_injection_text_is_stored_verbatim(api, text):
    x = api.lodge(depositor_name=text)
    assert api.get("rb.head", f"/instruments/{x['id']}").json()["depositor_name"] == text
    assert any(d["code"] == "NAME_MISMATCH" for d in x["discrepancies"])


def test_actions_with_wrong_shapes(api):
    x = api.lodge()
    for action in ["verify", "bank_reply", "accept", "reject", "rectify", "release_request", "release_decide",
                   "waive_discrepancy", "not_an_action", "__init__", "../admin"]:
        for body in [{}, {"officer_id": -1}, {"found": "maybe"}, {"bank_values": "x"}, {"remarks": "a" * 20000},
                     {"corrections": {"status": "RELEASED"}}, {"corrections": {"amount": "x"}}, {"checklist": []}]:
            r = api.post("rb.head", f"/instruments/{x['id']}/actions/{action}", body)
            assert r.status_code < 500, (action, body, r.status_code, r.text[:200])
    assert api.get("rb.head", f"/instruments/{x['id']}").json()["status"] != "RELEASED"


def test_uploads_are_checked(api):
    x = api.lodge()
    for name, data in [("a.exe", b"MZ\x90\x00"), ("a.pdf", b"not a pdf"), ("a.pdf", b"%PDF-" + b"0" * (11 << 20))]:
        r = api.client.post(f"/api/instruments/{x['id']}/documents", headers=api.h("rb.head"),
                            files={"file": (name, data, "application/pdf")})
        assert r.status_code in (413, 415, 422), (name, r.status_code)


def test_expired_and_forged_tokens(api):
    for h in ["", "Bearer", "Bearer x.y.z", "Bearer " + "a" * 2000, "Bearer eyJhbGciOiJub25lIn0.eyJzdWIiOiIxIn0.",
              "Basic YWRtaW46YWRtaW4="]:
        assert api.client.get("/api/instruments", headers={"Authorization": h}).status_code == 401


def test_double_accept_only_records_once(api):
    b = api.body()
    api.bank_knows(b)
    x = api.lodge(body=b)
    api.act("rb.head", x["id"], "verify")
    first = api.post("rb.head", f"/instruments/{x['id']}/actions/accept", {})
    second = api.post("rb.head", f"/instruments/{x['id']}/actions/accept", {})
    assert first.status_code == 200 and second.status_code == 409
    events = api.get("rb.head", f"/instruments/{x['id']}").json()["events"]
    assert sum(e["event_type"] == "ACCEPTED" for e in events) == 1


def test_exports_do_not_run_formulas(api):
    payload = '=HYPERLINK("http://evil.example","Click")'
    x = api.lodge(depositor_name=payload)
    h = api.h("rb.head")
    csv_body = api.client.get("/api/instruments/export", params={"format": "csv", "q": x["ref_no"]},
                              headers=h).content.decode("utf-8-sig")
    assert payload not in csv_body.replace("'" + payload.replace('"', '""'), "")
    from openpyxl import load_workbook

    xlsx = api.client.get("/api/instruments/export", params={"format": "xlsx", "q": x["ref_no"]}, headers=h).content
    cells = [c.value for row in load_workbook(io.BytesIO(xlsx)).active.iter_rows() for c in row]
    assert payload not in cells
